#!/usr/bin/env python3
"""
Build a UTM table (base URL / UTM parameters / full tagged URL) from a
campaign spec and a list of destination URLs.

This handles the deterministic, error-prone part of UTM tagging by hand:
- lowercasing and hyphenating values
- correctly appending query params to a URL that may already have some
  (using '?' vs '&' correctly, and not clobbering existing params)
- passing ValueTrack-style {placeholders} through untouched, braces and case
  intact, so ad-platform tokens still resolve at click time
- producing a clean three-column table as both CSV and a simple XLSX

For a paid-search tracking template, pass "{lpurl}" as the url and
"{keyword}" as utm_term; both survive normalization verbatim and the
resulting "Full Tagged URL" is the string to paste into the platform.

Usage:
    python build_utm_table.py campaign_spec.json output_basename

campaign_spec.json shape:
{
  "utm_source": "linkedin",
  "utm_medium": "paid-social",
  "utm_campaign": "2026-q3-product-launch",
  "utm_content": null,          # optional, string or null
  "utm_term": null,             # optional, string or null
  "urls": [
    {"url": "https://knak.com/product", "utm_content": null, "utm_term": null},
    {"url": "https://knak.com/pricing", "utm_content": "pricing-cta", "utm_term": null}
  ]
}

Per-URL utm_content/utm_term override the top-level value for that row when set.
"""

import csv
import json
import re
import sys
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

# --- dependency bootstrap -------------------------------------------------
# Cowork hands every session a fresh container, so this skill's packages may not
# be installed yet and no setup step will have survived from last time. Install
# them quietly here. Failure is not fatal: the script carries on with whatever it
# can still do, and says so where the result is affected.
import os as _os
import sys as _sys

_sys.path.insert(0, _os.path.abspath(_os.path.join(
    _os.path.dirname(_os.path.abspath(__file__)), "..", "..", "..", "scripts")))
try:
    from ensure_deps import ensure as _ensure
    _ensure("utm-generator")
except Exception:
    pass
# --------------------------------------------------------------------------

FIELDS = ["utm_source", "utm_medium", "utm_campaign", "utm_content", "utm_term"]


PLACEHOLDER_RE = re.compile(r"\{[^{}]*\}")


def _normalize_plain(value):
    """Lowercase, replace spaces/underscores with hyphens, strip illegal chars."""
    value = value.lower()
    value = re.sub(r"[\s_]+", "-", value)
    value = re.sub(r"[^a-z0-9\-\.]", "", value)
    return value


def normalize(value):
    """Normalize a UTM value, passing ValueTrack-style {placeholders} through untouched.

    Ad platforms substitute tokens like {keyword}, {matchtype} or {lpurl} at click
    time, so the braces have to survive into the final URL and the text inside them
    must not be case-folded — some tokens are case-sensitive, and a lowercased or
    brace-stripped token is never substituted, it just arrives in the report as
    literal text. Everything outside the braces is still normalized as usual, so a
    mixed value like 'Brand_{keyword}' becomes 'brand-{keyword}'.
    """
    if value is None:
        return None
    value = str(value).strip()

    segments = PLACEHOLDER_RE.split(value)
    placeholders = PLACEHOLDER_RE.findall(value)

    out = []
    for i, segment in enumerate(segments):
        out.append(_normalize_plain(segment))
        if i < len(placeholders):
            out.append(placeholders[i])
    result = "".join(out)

    # Collapse hyphen runs introduced by normalization. Safe over the whole string:
    # hyphens are never a placeholder delimiter, and strip('-') cannot eat a
    # leading '{' or trailing '}'.
    result = re.sub(r"-{2,}", "-", result).strip("-")
    return result


def build_utm_params(campaign, row):
    params = {}
    for field in FIELDS:
        # per-row override takes precedence over the campaign-level default
        value = row.get(field, None)
        if value is None:
            value = campaign.get(field, None)
        normalized = normalize(value)
        if normalized:
            params[field] = normalized
    return params


def append_params_to_url(base_url, params):
    parts = urlsplit(base_url)
    existing = dict(parse_qsl(parts.query, keep_blank_values=True))
    existing.update(params)  # UTM params win on key collision
    # safe="{}" keeps ValueTrack braces literal — a percent-encoded %7Bkeyword%7D
    # is not recognized by the ad platform and never gets substituted.
    new_query = urlencode(existing, safe="{}")
    return urlunsplit((parts.scheme, parts.netloc, parts.path, new_query, parts.fragment))


def utm_query_string(params):
    ordered = {k: params[k] for k in FIELDS if k in params}
    return urlencode(ordered, safe="{}")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    spec_path, out_base = sys.argv[1], sys.argv[2]
    with open(spec_path, "r") as f:
        campaign = json.load(f)

    urls = campaign.get("urls", [])
    if not urls:
        print("No URLs found in campaign spec under 'urls'.")
        sys.exit(1)

    rows = []
    warnings = []
    for i, entry in enumerate(urls):
        base_url = entry["url"].strip()
        params = build_utm_params(campaign, entry)

        missing_required = [f for f in ("utm_source", "utm_medium", "utm_campaign") if f not in params]
        if missing_required:
            warnings.append(f"Row {i+1} ({base_url}): missing required field(s) {missing_required}")

        query_string = utm_query_string(params)
        tagged_url = append_params_to_url(base_url, params)

        rows.append({
            "Base URL": base_url,
            "UTM Parameters": query_string,
            "Full Tagged URL": tagged_url,
        })

    csv_path = f"{out_base}.csv"
    with open(csv_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["Base URL", "UTM Parameters", "Full Tagged URL"])
        writer.writeheader()
        writer.writerows(rows)

    # Try to also produce an .xlsx if openpyxl is available; CSV always works regardless.
    xlsx_path = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "UTM Table"
        headers = ["Base URL", "UTM Parameters", "Full Tagged URL"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row["Base URL"], row["UTM Parameters"], row["Full Tagged URL"]])
        for col_idx, header in enumerate(headers, start=1):
            max_len = max([len(header)] + [len(str(r[header])) for r in rows]) if rows else len(header)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 80)
        xlsx_path = f"{out_base}.xlsx"
        wb.save(xlsx_path)
    except ImportError:
        pass

    print(json.dumps({
        "csv_path": csv_path,
        "xlsx_path": xlsx_path,
        "row_count": len(rows),
        "warnings": warnings,
    }, indent=2))


if __name__ == "__main__":
    main()
